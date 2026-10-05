#!/usr/bin/env python3
"""propx/ 폴더의 PropX(부동산114)(부동산114) '지역별 단지정보' 엑셀을 읽어 단지 시세 마스터를 만든다.

- 파일: propx/*.xlsx (PropX 엑셀저장 결과. 여러 지역 파일을 그대로 넣으면 됨. 같은 지역의 새 파일이 오면 최신 파일 우선)
- 열 이름은 PropX 화면 그대로 인식 (구분·도시·구시군·읍면동·아파트·전용면적·세대수·총세대수·입주년월·매매평균·매매상한·매매하한·전세평균·매매변동률·전세변동률·지하철·임대)
- 전용 49~85.5㎡만 사용, 49/59/74/84 로 묶어 세대수 가중 평균
- 결과: data/propx.json  { generated_at, files, complexes:[{key,name,sgg,umd,area,...}] }
"""
import json, re, sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "propx"
OUT = ROOT / "data" / "propx.json"
cfg = json.loads((ROOT / "config" / "targets.json").read_text(encoding="utf-8"))
disc = cfg.get("discovery", {})
AMIN, AMAX = disc.get("exclu_min", 49), disc.get("exclu_max", 85.5)
BUCKETS = disc.get("area_buckets", [49, 59, 74, 84])

COLS = {
    "gb": ["구분"], "city": ["도시"], "sgg": ["구시군", "시군구"], "umd": ["읍면동"],
    "code": ["아파트코드", "단지코드"], "name": ["아파트", "단지명", "아파트명"],
    "supply": ["분양면적"], "py": ["평형"], "area": ["전용면적"],
    "danji": ["총세대수"], "sedae": ["세대수"], "floors": ["총 층수", "총층수"],
    "bun_dt": ["분양년월"], "ibju": ["입주년월"],
    "mm": ["매매평균"], "mm_h": ["매매상한"], "mm_l": ["매매하한"], "js": ["전세평균"],
    "mm_chg": ["매매변동률"], "js_chg": ["전세변동률"], "subway": ["지하철"],
    "imdae": ["임대"], "japt": ["재건축"], "con": ["건설사"],
}


def nh(s):
    return re.sub(r"[\s()㎡%만원]", "", str(s or ""))


def map_cols(header):
    m = {}
    H = [nh(h) for h in header]
    for key, cands in COLS.items():
        for i, h in enumerate(H):
            if i in m.values():
                continue
            if any(h == nh(c) or (nh(c) and h.startswith(nh(c)) and key not in ("sedae",)) for c in cands) or (key == "sedae" and h == "세대수"):
                m[key] = i
                break
    return m


def num(v):
    if v in (None, "", "-"):
        return None
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return None


def read(path):
    import openpyxl, warnings
    warnings.simplefilter("ignore")
    wb = openpyxl.load_workbook(path, data_only=True)   # SheetJS 파일은 read_only 모드에서 행이 안 읽힘
    ws = wb[wb.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    # PropX 엑셀은 2행 헤더(그룹 '정보' + 세부 열). 두 행을 합쳐 하나의 헤더로 만든다
    hi = None
    for i in range(min(5, len(rows) - 1)):
        merged = [(rows[i + 1][j] if (rows[i][j] in (None, "", "정보")) and j < len(rows[i + 1]) and rows[i + 1][j] not in (None, "") else rows[i][j]) for j in range(len(rows[i]))]
        if "name" in map_cols(merged) and "area" in map_cols(merged):
            hi, header = i + 1, merged
            break
    if hi is None:
        hi = next((i for i, r in enumerate(rows[:5]) if "name" in map_cols(r) and "area" in map_cols(r)), None)
        if hi is None:
            print(f"::warning::{path.name}: 헤더를 찾지 못했습니다. 건너뜁니다.")
            return []
        header = rows[hi]
    m = map_cols(header)
    out = []
    for r in rows[hi + 1:]:
        g = lambda k: r[m[k]] if k in m and m[k] < len(r) else None
        if not g("name"):
            continue
        out.append({k: g(k) for k in COLS})
    return out


def bucket(a):
    if a is None or not (AMIN <= a <= AMAX):
        return None
    return min(BUCKETS, key=lambda b: abs(b - a))


def main():
    files = sorted(SRC.glob("*.xlsx"), key=lambda p: p.stat().st_mtime)
    if not files:
        OUT.write_text(json.dumps({"complexes": [], "files": []}, ensure_ascii=False), encoding="utf-8")
        print("propx/ 에 파일이 없습니다.")
        return
    # 같은 (구시군) 의 새 파일이 오면 이전 파일 행은 버림
    by_sgg = {}
    for f in files:
        rows = read(f)
        for sgg in {r["sgg"] for r in rows}:
            by_sgg[sgg] = (f.name, [r for r in rows if r["sgg"] == sgg])
    agg = defaultdict(lambda: {"w": 0, "mm": 0, "mm_h": 0, "mm_l": 0, "js": 0, "jsw": 0, "mm_chg": 0, "js_chg": 0, "jcw": 0, "sedae": 0, "n": 0, "rows": []})
    for sgg, (fname, rows) in by_sgg.items():
        for r in rows:
            b = bucket(num(r["area"]))
            if b is None:
                continue
            key = f"{r['sgg']}|{r['umd']}|{r['name']}|{b}"
            a = agg[key]
            w = num(r["sedae"]) or 1
            a["sedae"] += w; a["n"] += 1
            a["meta"] = {"sgg": r["sgg"], "umd": r["umd"], "name": str(r["name"]).strip(), "area": b, "code": r["code"],
                         "danji": num(r["danji"]), "ibju": r["ibju"], "subway": r["subway"], "gb": r["gb"], "imdae": r["imdae"], "file": fname,
                         "floors": r["floors"]}
            if num(r["mm"]):
                a["w"] += w; a["mm"] += num(r["mm"]) * w
                a["mm_h"] += (num(r["mm_h"]) or num(r["mm"])) * w; a["mm_l"] += (num(r["mm_l"]) or num(r["mm"])) * w
                if num(r["mm_chg"]) is not None:
                    a["mm_chg"] += num(r["mm_chg"]) * w
            if num(r["js"]):
                a["jsw"] += w; a["js"] += num(r["js"]) * w
                if num(r["js_chg"]) is not None:
                    a["jcw"] += w; a["js_chg"] += num(r["js_chg"]) * w
    out = []
    for key, a in agg.items():
        m = a["meta"]
        mm = a["mm"] / a["w"] / 10000 if a["w"] else None
        js = a["js"] / a["jsw"] / 10000 if a["jsw"] else None
        rec = {
            **m, "key": key, "sedae": int(a["sedae"]), "types": a["n"],
            "mm": round(mm, 3) if mm else None,
            "mm_h": round(a["mm_h"] / a["w"] / 10000, 3) if a["w"] else None,
            "mm_l": round(a["mm_l"] / a["w"] / 10000, 3) if a["w"] else None,
            "js": round(js, 3) if js else None,
            "mm_chg": round(a["mm_chg"] / a["w"], 1) if a["w"] else None,
            "js_chg": round(a["js_chg"] / a["jcw"], 1) if a["jcw"] else None,
        }
        if mm and js:
            rec["jeonse_ratio"] = round(100 * js / mm, 1)
            rec["gap"] = round(mm - js, 2)
            rec["required"] = round(mm - js + mm * cfg["acquisition_tax_rate"], 2)
        out.append(rec)
    out.sort(key=lambda x: (x["sgg"], x["umd"], x["name"], x["area"]))
    OUT.write_text(json.dumps({
        "generated_at": datetime.now().strftime("%Y-%m-%d"),
        "files": [{"sgg": s, "file": f} for s, (f, _) in by_sgg.items()],
        "complexes": out,
    }, ensure_ascii=False, indent=0), encoding="utf-8")
    print(f"propx.json: {len(out)}개 단지·평형 · 지역 {sorted(by_sgg)}")


if __name__ == "__main__":
    main()
# v1.2
