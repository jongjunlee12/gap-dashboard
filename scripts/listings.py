#!/usr/bin/env python3
"""listings/ 폴더의 매물 파일(xlsx·csv)을 읽어 정규화하고, 직전 파일과 비교합니다.

- 파일명 날짜순으로 정렬 (예: 2026-10-05.xlsx). 날짜가 없으면 수정시각 사용
- 최신 파일 = 현재 매물, 그 직전 파일 = 비교 기준
- 열 이름은 유사어로 자동 인식 (단지/아파트, 전용/면적/평형, 동, 층, 호가/매매가/가격, 전세/보증금, 등록일, 출처)
- 결과: data/listings.json
"""
import json, re, sys, csv
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
LIST = ROOT / "listings"
OUT = ROOT / "data" / "listings.json"

COLS = {
    "complex": ["단지", "아파트", "단지명", "아파트명", "물건명", "complex", "name"],
    "area": ["전용", "전용면적", "면적", "평형", "공급", "area", "㎡"],
    "dong": ["동", "동명", "건물동", "dong"],
    "floor": ["층", "해당층", "층수", "floor"],
    "price": ["호가", "매매가", "매매가격", "가격", "매도가", "price", "매매"],
    "jeonse": ["전세", "전세가", "보증금", "전세보증금", "jeonse", "deposit"],
    "type": ["거래유형", "유형", "구분", "type"],
    "date": ["등록일", "확인일", "날짜", "일자", "date"],
    "source": ["출처", "사이트", "source", "중개사", "부동산"],
    "note": ["비고", "메모", "특징", "note", "설명"],
}


def norm_header(h):
    return re.sub(r"[\s_()\[\]㎡m²]", "", str(h or "")).lower()


def map_columns(headers):
    m = {}
    nh = [norm_header(h) for h in headers]
    for key, cands in COLS.items():
        for i, h in enumerate(nh):
            if i in m.values():
                continue
            if any(norm_header(c) == h or (len(h) >= 2 and norm_header(c) in h) for c in cands):
                m[key] = i
                break
    return m


def to_eok(v):
    """'9억 2,000' / '92,000'(만원) / '9.2' / 920000000 → 억 단위 float"""
    if v is None:
        return None
    s = str(v).replace(",", "").replace(" ", "").strip()
    if not s or s in ("-", "nan", "None"):
        return None
    mm = re.match(r"^(\d+(?:\.\d+)?)억(?:(\d+))?(?:만)?$", s)
    if mm:
        eok = float(mm.group(1))
        if mm.group(2):
            eok += float(mm.group(2)) / 10000
        return round(eok, 3)
    try:
        x = float(re.sub(r"[^\d.]", "", s))
    except ValueError:
        return None
    if x > 1e7:            # 원
        return round(x / 1e8, 3)
    if x > 1000:           # 만원
        return round(x / 10000, 3)
    return round(x, 3)     # 억


def to_area(v):
    if v is None:
        return None
    s = str(v)
    m = re.search(r"(\d+(?:\.\d+)?)", s)
    if not m:
        return None
    x = float(m.group(1))
    if "평" in s and x < 60:
        x = x * 3.3058
    return int(round(x))


def to_floor(v):
    if v is None:
        return None
    m = re.search(r"(\d+)", str(v))
    return int(m.group(1)) if m else None


def read_rows(path: Path):
    if path.suffix.lower() == ".csv":
        with open(path, encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f))
    else:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        return []
    # 헤더 행: 단지/아파트 비슷한 열이 있는 첫 행
    hi = 0
    for i, r in enumerate(rows[:10]):
        if "complex" in map_columns(r):
            hi = i
            break
    m = map_columns(rows[hi])
    out = []
    for r in rows[hi + 1:]:
        g = lambda k: r[m[k]] if k in m and m[k] < len(r) else None
        if not g("complex"):
            continue
        t = str(g("type") or "")
        price = to_eok(g("price"))
        jeonse = to_eok(g("jeonse"))
        if "전세" in t and price and not jeonse:
            jeonse, price = price, None
        out.append({
            "complex_raw": str(g("complex")).strip(),
            "area": to_area(g("area")),
            "dong": (str(g("dong")).strip() if g("dong") not in (None, "") else None),
            "floor": to_floor(g("floor")),
            "price": price,
            "jeonse": jeonse,
            "type": "전세" if ("전세" in t or (jeonse and not price)) else "매매",
            "date": (str(g("date"))[:10] if g("date") else None),
            "source": (str(g("source")).strip() if g("source") else None),
            "note": (str(g("note")).strip() if g("note") else None),
        })
    return out


def file_date(p: Path):
    m = re.search(r"(20\d{2})[-._]?(\d{2})[-._]?(\d{2})", p.stem)
    if m:
        return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    return datetime.fromtimestamp(p.stat().st_mtime).strftime("%Y-%m-%d")


def key(it):
    return (it["complex_raw"], it["area"], it["dong"], it["floor"], it["type"])


def main():
    files = sorted(
        [p for p in LIST.glob("*") if p.suffix.lower() in (".xlsx", ".xlsm", ".csv") and not p.name.startswith("_")],
        key=lambda p: (file_date(p), p.name),
    )
    if not files:
        OUT.write_text(json.dumps({"items": [], "files": []}, ensure_ascii=False, indent=1), encoding="utf-8")
        print("listings/ 에 파일이 없습니다.")
        return
    cur = read_rows(files[-1])
    prev = read_rows(files[-2]) if len(files) > 1 else []
    pk = {key(i): i for i in prev}
    ck = {key(i): i for i in cur}
    for it in cur:
        k = key(it)
        it["status"] = "new" if k not in pk else "kept"
        if k in pk:
            p0, p1 = pk[k].get("price"), it.get("price")
            if p0 and p1 and p1 < p0:
                it["status"] = "reduced"
                it["prev_price"] = p0
    removed = [dict(pk[k], status="removed") for k in pk if k not in ck]
    result = {
        "date": file_date(files[-1]),
        "file": files[-1].name,
        "prev_file": files[-2].name if len(files) > 1 else None,
        "files": [p.name for p in files],
        "items": cur,
        "removed": removed,
        "summary": {
            "total": len(cur),
            "new": sum(i["status"] == "new" for i in cur),
            "reduced": sum(i["status"] == "reduced" for i in cur),
            "removed": len(removed),
        },
    }
    OUT.write_text(json.dumps(result, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{files[-1].name}: {len(cur)}건 (신규 {result['summary']['new']} · 인하 {result['summary']['reduced']} · 사라짐 {len(removed)})")


if __name__ == "__main__":
    main()
